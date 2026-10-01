#!/usr/bin/env python3
"""Generate PostgreSQL gift-catalog SQL from the gift workbook template."""

from __future__ import annotations

import argparse
import json
import sys
import unicodedata
import uuid
from dataclasses import dataclass, replace
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


GIFT_HEADERS = {
    "importar?",
    "codigo",
    "categoria",
    "nome",
    "descricao",
    "tipo de presente",
    "valor total / referencia (r$)",
    "quantidade de cotas",
    "modo de compra",
    "url da imagem",
    "url de pagamento com cartao",
}

OPTION_HEADERS = {
    "codigo do presente",
    "ordem",
    "tipo da opcao",
    "loja",
    "url do produto",
    "observacoes",
}

TRUE_VALUES = {"sim", "s", "yes", "true", "1"}
FALSE_VALUES = {"nao", "n", "no", "false", "0"}

GIFT_TYPES = {
    "individual": "single",
    "cotas": "quota",
}

PURCHASE_MODES = {
    "dinheiro / pix / cartao": "money",
    "compra externa": "external",
    "hibrido": "hybrid",
}

OPTION_TYPES = {
    "online": "online",
    "loja fisica": "physical",
}


class WorkbookValidationError(ValueError):
    """Raised when the workbook does not follow the gift import layout."""


@dataclass(frozen=True)
class PurchaseOption:
    source_row: int
    order: int
    option_type: str
    store: str
    url: str
    notes: str

    def as_json(self) -> dict[str, str]:
        return {
            "type": self.option_type,
            "store": self.store,
            "url": self.url if self.option_type == "online" else "",
            "notes": self.notes,
        }


@dataclass(frozen=True)
class Gift:
    source_row: int
    workbook_code: str
    gift_id: uuid.UUID
    category: str
    name: str
    description: str
    price: Decimal | None
    image_url: str
    card_payment_url: str
    gift_type: str
    quota_count: int | None
    purchase_mode: str
    purchase_options: tuple[PurchaseOption, ...]


def normalize(value: object) -> str:
    """Normalize labels for accent-, spacing-, and case-insensitive matching."""
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(text.strip().casefold().split())


def display_text(value: object) -> str:
    return " ".join(("" if value is None else str(value)).strip().split())


def header_map(sheet: Worksheet) -> dict[str, int]:
    return {
        normalize(sheet.cell(row=1, column=column).value): column
        for column in range(1, sheet.max_column + 1)
        if display_text(sheet.cell(row=1, column=column).value)
    }


def select_sheet(
    workbook,
    required_headers: set[str],
    requested_sheet: str | None,
    purpose: str,
) -> Worksheet:
    if requested_sheet:
        if requested_sheet not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise WorkbookValidationError(
                f"A aba {requested_sheet!r} não existe. Abas disponíveis: {available}."
            )
        sheet = workbook[requested_sheet]
        missing = required_headers - set(header_map(sheet))
        if missing:
            raise WorkbookValidationError(
                f"A aba {requested_sheet!r} não contém os cabeçalhos de {purpose}: "
                + ", ".join(sorted(missing))
                + "."
            )
        return sheet

    matches = [
        sheet
        for sheet in workbook.worksheets
        if required_headers <= set(header_map(sheet))
    ]
    if not matches:
        raise WorkbookValidationError(
            f"Nenhuma aba contém todos os cabeçalhos de {purpose}: "
            + ", ".join(sorted(required_headers))
            + "."
        )
    if len(matches) > 1:
        names = ", ".join(sheet.title for sheet in matches)
        raise WorkbookValidationError(
            f"Mais de uma aba possui a estrutura de {purpose} ({names}). "
            "Informe o nome da aba explicitamente."
        )
    return matches[0]


def parse_import_flag(value: object, *, row: int) -> bool:
    key = normalize(value)
    if key in TRUE_VALUES:
        return True
    if key in FALSE_VALUES:
        return False
    raise WorkbookValidationError(
        f"Linha {row} da aba de presentes: valor inválido em IMPORTAR?: "
        f"{display_text(value)!r}. Use Sim ou Não."
    )


def parse_decimal(value: object, *, row: int, field: str) -> Decimal | None:
    if value is None or display_text(value) == "":
        return None
    if isinstance(value, bool):
        raise WorkbookValidationError(
            f"Linha {row}: {field} deve ser um número maior que zero."
        )
    try:
        if isinstance(value, (int, float, Decimal)):
            number = Decimal(str(value))
        else:
            text = display_text(value).replace("R$", "").replace(" ", "")
            if "," in text:
                text = text.replace(".", "").replace(",", ".")
            number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise WorkbookValidationError(
            f"Linha {row}: valor inválido em {field}: {display_text(value)!r}."
        ) from None
    if not number.is_finite() or number <= 0:
        raise WorkbookValidationError(
            f"Linha {row}: {field} deve ser maior que zero."
        )
    return number


def parse_positive_integer(value: object, *, row: int, field: str) -> int | None:
    number = parse_decimal(value, row=row, field=field)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise WorkbookValidationError(
            f"Linha {row}: {field} deve ser um número inteiro."
        )
    return int(number)


def is_http_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme.casefold() in {"http", "https"} and bool(parsed.netloc)


def validate_length(value: str, limit: int, *, row: int, field: str) -> None:
    if len(value) > limit:
        raise WorkbookValidationError(
            f"Linha {row}: {field} excede o limite de {limit} caracteres."
        )


def row_has_content(sheet: Worksheet, row: int, columns: Iterable[int]) -> bool:
    return any(display_text(sheet.cell(row=row, column=column).value) for column in columns)


def parse_gift_rows(sheet: Worksheet) -> list[Gift]:
    columns = header_map(sheet)
    relevant_columns = [columns[header] for header in GIFT_HEADERS]
    gifts: list[Gift] = []
    used_codes: set[str] = set()

    for row in range(2, sheet.max_row + 1):
        if not row_has_content(sheet, row, relevant_columns):
            continue

        raw_import = sheet.cell(row=row, column=columns["importar?"]).value
        if not display_text(raw_import):
            raise WorkbookValidationError(
                f"Linha {row} da aba de presentes: preencha IMPORTAR? com Sim ou Não."
            )
        if not parse_import_flag(raw_import, row=row):
            continue

        code = display_text(sheet.cell(row=row, column=columns["codigo"]).value)
        category = display_text(sheet.cell(row=row, column=columns["categoria"]).value)
        name = display_text(sheet.cell(row=row, column=columns["nome"]).value)
        description = display_text(
            sheet.cell(row=row, column=columns["descricao"]).value
        )
        image_url = display_text(
            sheet.cell(row=row, column=columns["url da imagem"]).value
        )
        card_payment_url = display_text(
            sheet.cell(row=row, column=columns["url de pagamento com cartao"]).value
        )

        missing = [
            label
            for label, value in (("CÓDIGO", code), ("CATEGORIA", category), ("NOME", name))
            if not value
        ]
        if missing:
            raise WorkbookValidationError(
                f"Linha {row}: campos obrigatórios vazios: {', '.join(missing)}."
            )
        normalized_code = normalize(code)
        if normalized_code in used_codes:
            raise WorkbookValidationError(f"Linha {row}: código duplicado: {code!r}.")
        used_codes.add(normalized_code)

        raw_type = sheet.cell(row=row, column=columns["tipo de presente"]).value
        type_key = normalize(raw_type)
        if type_key not in GIFT_TYPES:
            raise WorkbookValidationError(
                f"Linha {row}: tipo de presente inválido: {display_text(raw_type)!r}."
            )
        gift_type = GIFT_TYPES[type_key]

        raw_mode = sheet.cell(row=row, column=columns["modo de compra"]).value
        mode_key = normalize(raw_mode)
        if mode_key not in PURCHASE_MODES:
            raise WorkbookValidationError(
                f"Linha {row}: modo de compra inválido: {display_text(raw_mode)!r}."
            )
        purchase_mode = PURCHASE_MODES[mode_key]

        price = parse_decimal(
            sheet.cell(row=row, column=columns["valor total / referencia (r$)"]).value,
            row=row,
            field="VALOR TOTAL / REFERÊNCIA (R$)",
        )
        quota_count = parse_positive_integer(
            sheet.cell(row=row, column=columns["quantidade de cotas"]).value,
            row=row,
            field="QUANTIDADE DE COTAS",
        )

        if gift_type == "quota":
            if price is None or quota_count is None:
                raise WorkbookValidationError(
                    f"Linha {row}: presente por Cotas exige valor total e quantidade de cotas."
                )
            if purchase_mode != "money":
                raise WorkbookValidationError(
                    f"Linha {row}: presente por Cotas deve usar "
                    "Dinheiro / PIX / Cartão."
                )
        else:
            if quota_count is not None:
                raise WorkbookValidationError(
                    f"Linha {row}: presente Individual não deve informar quantidade de cotas."
                )
            if purchase_mode in {"money", "hybrid"} and price is None:
                raise WorkbookValidationError(
                    f"Linha {row}: o modo de compra exige um valor de referência."
                )

        for field, value, limit in (
            ("CATEGORIA", category, 200),
            ("NOME", name, 300),
            ("DESCRIÇÃO", description, 4000),
            ("URL DA IMAGEM", image_url, 2000),
            ("URL DE PAGAMENTO COM CARTÃO", card_payment_url, 2000),
        ):
            validate_length(value, limit, row=row, field=field)

        if image_url and not is_http_url(image_url):
            raise WorkbookValidationError(f"Linha {row}: URL da imagem inválida.")
        if card_payment_url and not is_http_url(card_payment_url):
            raise WorkbookValidationError(
                f"Linha {row}: URL de pagamento com cartão inválida."
            )

        gifts.append(
            Gift(
                source_row=row,
                workbook_code=code,
                gift_id=uuid.uuid4(),
                category=category,
                name=name,
                description=description,
                price=price,
                image_url=image_url,
                card_payment_url=card_payment_url,
                gift_type=gift_type,
                quota_count=quota_count,
                purchase_mode=purchase_mode,
                purchase_options=(),
            )
        )

    if not gifts:
        raise WorkbookValidationError(
            "A aba de presentes não contém nenhuma linha marcada para importação."
        )
    return gifts


def skipped_gift_codes(sheet: Worksheet) -> set[str]:
    """Return codes explicitly excluded with IMPORTAR? = Não."""
    columns = header_map(sheet)
    skipped: set[str] = set()
    for row in range(2, sheet.max_row + 1):
        raw_import = sheet.cell(row=row, column=columns["importar?"]).value
        code = display_text(sheet.cell(row=row, column=columns["codigo"]).value)
        if code and normalize(raw_import) in FALSE_VALUES:
            skipped.add(normalize(code))
    return skipped


def parse_purchase_options(
    sheet: Worksheet,
    gifts: list[Gift],
    excluded_codes: set[str] | None = None,
) -> dict[str, tuple[PurchaseOption, ...]]:
    columns = header_map(sheet)
    relevant_columns = [columns[header] for header in OPTION_HEADERS]
    gifts_by_code = {normalize(gift.workbook_code): gift for gift in gifts}
    options: dict[str, list[PurchaseOption]] = {
        normalize(gift.workbook_code): [] for gift in gifts
    }
    used_orders: set[tuple[str, int]] = set()

    for row in range(2, sheet.max_row + 1):
        if not row_has_content(sheet, row, relevant_columns):
            continue

        code = display_text(
            sheet.cell(row=row, column=columns["codigo do presente"]).value
        )
        if not code:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: informe o CÓDIGO DO PRESENTE."
            )
        code_key = normalize(code)
        if code_key not in gifts_by_code:
            if excluded_codes and code_key in excluded_codes:
                continue
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: código não importado ou inexistente: {code!r}."
            )

        order = parse_positive_integer(
            sheet.cell(row=row, column=columns["ordem"]).value,
            row=row,
            field="ORDEM",
        )
        if order is None:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: informe a ORDEM."
            )
        if (code_key, order) in used_orders:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: ordem {order} duplicada para {code!r}."
            )
        used_orders.add((code_key, order))

        raw_type = sheet.cell(row=row, column=columns["tipo da opcao"]).value
        type_key = normalize(raw_type)
        if type_key not in OPTION_TYPES:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: tipo inválido: "
                f"{display_text(raw_type)!r}."
            )
        option_type = OPTION_TYPES[type_key]
        store = display_text(sheet.cell(row=row, column=columns["loja"]).value)
        url = display_text(sheet.cell(row=row, column=columns["url do produto"]).value)
        notes = display_text(sheet.cell(row=row, column=columns["observacoes"]).value)

        if not store:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: informe a LOJA."
            )
        if option_type == "online" and not is_http_url(url):
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: opção Online exige uma URL válida."
            )
        if option_type == "physical" and url:
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: Loja física não deve possuir URL."
            )
        validate_length(store, 200, row=row, field="LOJA")
        validate_length(url, 2000, row=row, field="URL DO PRODUTO")
        validate_length(notes, 1000, row=row, field="OBSERVAÇÕES")

        gift = gifts_by_code[code_key]
        if gift.gift_type == "quota" or gift.purchase_mode == "money":
            raise WorkbookValidationError(
                f"Linha {row} da aba de opções: {code!r} usa somente pagamento em dinheiro."
            )

        options[code_key].append(
            PurchaseOption(
                source_row=row,
                order=order,
                option_type=option_type,
                store=store,
                url=url,
                notes=notes,
            )
        )

    return {
        code: tuple(sorted(items, key=lambda option: option.order))
        for code, items in options.items()
    }


def attach_purchase_options(
    gifts: list[Gift], options_by_code: dict[str, tuple[PurchaseOption, ...]]
) -> list[Gift]:
    return [
        replace(
            gift,
            purchase_options=options_by_code[normalize(gift.workbook_code)],
        )
        for gift in gifts
    ]


def sql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def sql_nullable_string(value: str) -> str:
    return sql_string(value) if value else "null"


def decimal_sql(value: Decimal | None) -> str:
    if value is None:
        return "null"
    return format(value, "f")


def options_sql(options: tuple[PurchaseOption, ...]) -> str:
    payload = json.dumps(
        [option.as_json() for option in options],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return f"{sql_string(payload)}::jsonb"


def build_sql(
    gifts: Iterable[Gift],
    source: Path,
    gifts_sheet_name: str,
    options_sheet_name: str,
) -> str:
    gift_list = list(gifts)
    rows: list[str] = []
    for index, gift in enumerate(gift_list):
        quota_value = (
            f"({decimal_sql(gift.price)}::numeric / {gift.quota_count})"
            if gift.gift_type == "quota"
            else "null"
        )
        values = (
            sql_string(str(gift.gift_id)),
            sql_string(gift.category),
            sql_string(gift.name),
            sql_nullable_string(gift.description),
            decimal_sql(gift.price),
            sql_nullable_string(gift.image_url),
            sql_string("Disponível"),
            "null",
            sql_nullable_string(gift.card_payment_url),
            "null",
            "null",
            sql_string(gift.purchase_mode),
            options_sql(gift.purchase_options),
            "null",
            "null",
            sql_string(gift.gift_type),
            str(gift.quota_count) if gift.quota_count is not None else "null",
            quota_value,
        )
        separator = "," if index < len(gift_list) - 1 else ";"
        rows.append(
            "  ("
            + ", ".join(values)
            + f"){separator} -- Excel: presente {gift.workbook_code}, linha {gift.source_row}"
        )

    return f"""-- Generated by scripts/generate_gift_import_sql.py
-- Source workbook: {source.name}
-- Gift sheet: {gifts_sheet_name}
-- Purchase options sheet: {options_sheet_name}
-- Gifts: {len(gift_list)}
-- Review this file before running it in the Supabase SQL Editor.
-- The generated UUIDs make rerunning this same SQL fail safely on the primary key.

begin;

insert into public.gifts (
  id,
  category,
  name,
  description,
  price,
  image_url,
  status,
  payment_status,
  card_payment_url,
  card_payment_provider,
  card_payment_reference,
  purchase_mode,
  external_purchase_options,
  selected_purchase_method,
  selected_purchase_details,
  gift_type,
  quota_count,
  quota_value
)
values
{chr(10).join(rows)}

commit;
"""


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Gera SQL do catálogo de presentes a partir da planilha modelo."
    )
    parser.add_argument("workbook", type=Path, help="Arquivo .xlsx de entrada.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Arquivo .sql de saída. Padrão: mesmo nome e pasta do Excel.",
    )
    parser.add_argument(
        "--gifts-sheet",
        help="Nome exato da aba de presentes. Por padrão, detectada pelos cabeçalhos.",
    )
    parser.add_argument(
        "--options-sheet",
        help="Nome exato da aba de opções. Por padrão, detectada pelos cabeçalhos.",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_argument_parser().parse_args(argv)
    workbook_path = args.workbook.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else workbook_path.with_suffix(".sql")
    )

    if workbook_path.suffix.casefold() != ".xlsx":
        print("Erro: o arquivo de entrada deve ter extensão .xlsx.", file=sys.stderr)
        return 2
    if not workbook_path.is_file():
        print(f"Erro: arquivo não encontrado: {workbook_path}", file=sys.stderr)
        return 2
    if output_path.suffix.casefold() != ".sql":
        print("Erro: o arquivo de saída deve ter extensão .sql.", file=sys.stderr)
        return 2
    if output_path == workbook_path:
        print("Erro: o arquivo de saída não pode substituir a planilha.", file=sys.stderr)
        return 2

    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        gifts_sheet = select_sheet(
            workbook, GIFT_HEADERS, args.gifts_sheet, "presentes"
        )
        options_sheet = select_sheet(
            workbook, OPTION_HEADERS, args.options_sheet, "opções de compra"
        )
        gifts = parse_gift_rows(gifts_sheet)
        options = parse_purchase_options(
            options_sheet,
            gifts,
            excluded_codes=skipped_gift_codes(gifts_sheet),
        )
        gifts = attach_purchase_options(gifts, options)
        sql = build_sql(
            gifts,
            workbook_path,
            gifts_sheet.title,
            options_sheet.title,
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(sql, encoding="utf-8", newline="\n")
    except (OSError, WorkbookValidationError) as error:
        print(f"Erro: {error}", file=sys.stderr)
        return 2

    quota_gifts = sum(gift.gift_type == "quota" for gift in gifts)
    option_count = sum(len(gift.purchase_options) for gift in gifts)
    print(f"Aba de presentes: {gifts_sheet.title}")
    print(f"Aba de opções: {options_sheet.title}")
    print(f"Presentes: {len(gifts)} ({quota_gifts} por cotas)")
    print(f"Opções de compra: {option_count}")
    print(f"SQL: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
