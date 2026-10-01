#!/usr/bin/env python3
"""Generate a PostgreSQL guest import from the wedding guest workbook.

The script understands invitation groups represented by vertically merged cells.
It writes SQL for the current ``public.guests`` schema used by this repository.
"""

from __future__ import annotations

import argparse
import json
import secrets
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


INVITE_CODE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
INVITE_CODE_LENGTH = 8

REQUIRED_HEADERS = {
    "tipo de convite",
    "nome",
    "papel no convite",
    "convidado de",
    "convite enviado",
    "save the date enviado",
}

INVITE_TYPES = {
    "individual": "individual",
    "casal": "couple",
}

ROLES = {
    "convidado": "guest",
    "membro": "member",
    "acompanhante": "companion",
}

GUEST_SIDES = {
    "noiva": "bride",
    "noivo": "groom",
    "ambos": "couple",
    "casal": "couple",
}

TRUE_VALUES = {"sim", "s", "yes", "true", "1"}
FALSE_VALUES = {"nao", "n", "no", "false", "0", ""}


class WorkbookValidationError(ValueError):
    """Raised when the workbook does not follow the expected guest layout."""


@dataclass(frozen=True)
class Invitation:
    source_rows: tuple[int, int]
    name: str
    invite_type: str
    couple_members: tuple[str, str] | None
    max_guests: int
    invite_sent: bool
    save_the_date_sent: bool
    guest_side: str
    invite_code: str


def normalize(value: object) -> str:
    """Normalize workbook labels for accent- and case-insensitive matching."""
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(text.strip().casefold().split())


def display_text(value: object) -> str:
    return " ".join(("" if value is None else str(value)).strip().split())


def header_map(sheet: Worksheet) -> dict[str, int]:
    return {
        normalize(sheet.cell(row=1, column=column).value): column
        for column in range(1, sheet.max_column + 1)
        if display_text(sheet.cell(row=1, column=column).value)
    }


def select_sheet(workbook, requested_sheet: str | None) -> Worksheet:
    if requested_sheet:
        if requested_sheet not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise WorkbookValidationError(
                f"A aba {requested_sheet!r} não existe. Abas disponíveis: {available}."
            )
        sheet = workbook[requested_sheet]
        missing = REQUIRED_HEADERS - set(header_map(sheet))
        if missing:
            raise WorkbookValidationError(
                f"A aba {requested_sheet!r} não contém os cabeçalhos: "
                + ", ".join(sorted(missing))
                + "."
            )
        return sheet

    matches = [
        sheet
        for sheet in workbook.worksheets
        if REQUIRED_HEADERS <= set(header_map(sheet))
    ]
    if not matches:
        raise WorkbookValidationError(
            "Nenhuma aba contém todos os cabeçalhos esperados: "
            + ", ".join(sorted(REQUIRED_HEADERS))
            + "."
        )
    if len(matches) > 1:
        names = ", ".join(sheet.title for sheet in matches)
        raise WorkbookValidationError(
            f"Mais de uma aba possui a estrutura esperada ({names}). Use --sheet."
        )
    return matches[0]


def merged_row_span(sheet: Worksheet, row: int, column: int) -> tuple[int, int]:
    for merged_range in sheet.merged_cells.ranges:
        if (
            merged_range.min_col <= column <= merged_range.max_col
            and merged_range.min_row <= row <= merged_range.max_row
        ):
            return merged_range.min_row, merged_range.max_row
    return row, row


def parse_boolean(value: object, *, row: int, field: str) -> bool:
    normalized = normalize(value)
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    raise WorkbookValidationError(
        f"Linha {row}: valor inválido em {field}: {display_text(value)!r}. "
        "Use Sim ou Não."
    )


def parse_invitations(sheet: Worksheet) -> list[Invitation]:
    columns = header_map(sheet)
    type_column = columns["tipo de convite"]
    name_column = columns["nome"]
    role_column = columns["papel no convite"]
    side_column = columns["convidado de"]
    invite_sent_column = columns["convite enviado"]
    save_sent_column = columns["save the date enviado"]

    invitations: list[Invitation] = []
    used_codes: set[str] = set()
    row = 2

    while row <= sheet.max_row:
        raw_type = sheet.cell(row=row, column=type_column).value
        raw_name = sheet.cell(row=row, column=name_column).value

        if not display_text(raw_type) and not display_text(raw_name):
            row += 1
            continue
        if not display_text(raw_type):
            raise WorkbookValidationError(
                f"Linha {row}: participante fora de um conjunto mesclado de convite."
            )

        group_start, group_end = merged_row_span(sheet, row, type_column)
        if group_start != row:
            raise WorkbookValidationError(
                f"Linha {row}: início inesperado dentro do convite das linhas "
                f"{group_start}-{group_end}."
            )

        invite_type_key = normalize(raw_type)
        if invite_type_key not in INVITE_TYPES:
            raise WorkbookValidationError(
                f"Linha {row}: tipo de convite inválido: {display_text(raw_type)!r}."
            )
        invite_type = INVITE_TYPES[invite_type_key]

        participants: list[tuple[int, str, str]] = []
        for participant_row in range(group_start, group_end + 1):
            participant_name = display_text(
                sheet.cell(row=participant_row, column=name_column).value
            )
            role_key = normalize(sheet.cell(row=participant_row, column=role_column).value)
            if not participant_name:
                raise WorkbookValidationError(
                    f"Linha {participant_row}: o nome do participante está vazio."
                )
            if role_key not in ROLES:
                raw_role = sheet.cell(row=participant_row, column=role_column).value
                raise WorkbookValidationError(
                    f"Linha {participant_row}: papel inválido: {display_text(raw_role)!r}."
                )
            participants.append((participant_row, participant_name, ROLES[role_key]))

        companions = [name for _, name, role in participants if role == "companion"]
        couple_members = [name for _, name, role in participants if role == "member"]
        primary_guests = [name for _, name, role in participants if role == "guest"]

        if invite_type == "individual":
            if len(primary_guests) != 1 or couple_members:
                raise WorkbookValidationError(
                    f"Linhas {group_start}-{group_end}: convite Individual deve ter "
                    "exatamente um Convidado e nenhum Membro."
                )
            invitation_name = primary_guests[0]
            member_pair = None
        else:
            if len(couple_members) != 2 or primary_guests:
                raise WorkbookValidationError(
                    f"Linhas {group_start}-{group_end}: convite Casal deve ter "
                    "exatamente dois Membros e nenhum Convidado principal."
                )
            invitation_name = f"{couple_members[0]} e {couple_members[1]}"
            member_pair = (couple_members[0], couple_members[1])

        side_key = normalize(sheet.cell(row=row, column=side_column).value)
        if side_key not in GUEST_SIDES:
            raw_side = sheet.cell(row=row, column=side_column).value
            raise WorkbookValidationError(
                f"Linha {row}: origem do convite inválida: {display_text(raw_side)!r}."
            )

        code = generate_unique_invite_code(used_codes)
        invitations.append(
            Invitation(
                source_rows=(group_start, group_end),
                name=invitation_name,
                invite_type=invite_type,
                couple_members=member_pair,
                max_guests=len(companions),
                invite_sent=parse_boolean(
                    sheet.cell(row=row, column=invite_sent_column).value,
                    row=row,
                    field="Convite Enviado",
                ),
                save_the_date_sent=parse_boolean(
                    sheet.cell(row=row, column=save_sent_column).value,
                    row=row,
                    field="Save the Date Enviado",
                ),
                guest_side=GUEST_SIDES[side_key],
                invite_code=code,
            )
        )
        row = group_end + 1

    if not invitations:
        raise WorkbookValidationError("A aba selecionada não contém convites.")
    return invitations


def generate_unique_invite_code(used_codes: set[str]) -> str:
    while True:
        code = "".join(
            secrets.choice(INVITE_CODE_ALPHABET) for _ in range(INVITE_CODE_LENGTH)
        )
        if code not in used_codes:
            used_codes.add(code)
            return code


def sql_string(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def sql_boolean(value: bool) -> str:
    return "true" if value else "false"


def sql_couple_members(members: tuple[str, str] | None) -> str:
    if members is None:
        return "null"
    payload = json.dumps(
        [{"name": members[0]}, {"name": members[1]}],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return f"{sql_string(payload)}::jsonb"


def build_sql(invitations: Iterable[Invitation], source: Path, sheet_name: str) -> str:
    invitation_list = list(invitations)
    rows = []
    for index, invitation in enumerate(invitation_list):
        values = (
            sql_string(invitation.name),
            sql_string(invitation.invite_code),
            str(invitation.max_guests),
            "false",
            sql_boolean(invitation.invite_sent),
            sql_boolean(invitation.save_the_date_sent),
            "true",
            "0",
            sql_string(invitation.invite_type),
            sql_couple_members(invitation.couple_members),
            sql_string(invitation.guest_side),
        )
        source_rows = (
            str(invitation.source_rows[0])
            if invitation.source_rows[0] == invitation.source_rows[1]
            else f"{invitation.source_rows[0]}-{invitation.source_rows[1]}"
        )
        separator = "," if index < len(invitation_list) - 1 else ";"
        rows.append(
            "  ("
            + ", ".join(values)
            + f"){separator} -- Excel: linhas {source_rows}"
        )

    values_sql = "\n".join(rows)
    return f"""-- Generated by scripts/generate_guest_import_sql.py
-- Source workbook: {source.name}
-- Source sheet: {sheet_name}
-- Invitations: {len(invitation_list)}
-- Review this file before running it in the Supabase SQL Editor.
-- The transaction fails safely if an invitation code already exists.

begin;

insert into public.guests (
  name,
  invite_code,
  max_guests,
  confirmed,
  invite_sent,
  save_the_date_sent,
  active,
  access_count,
  invite_type,
  couple_members,
  guest_side
)
values
{values_sql}

commit;
"""


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Gera SQL de convidados a partir da planilha do casamento."
    )
    parser.add_argument("workbook", type=Path, help="Arquivo .xlsx de entrada.")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Arquivo .sql de saída. Padrão: mesmo nome e pasta do Excel.",
    )
    parser.add_argument(
        "--sheet",
        help="Nome exato da aba. Sem esta opção, a aba é detectada pelos cabeçalhos.",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_argument_parser().parse_args(argv)
    workbook_path = args.workbook.expanduser().resolve()
    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else workbook_path.with_suffix(".sql")
    )

    if workbook_path.suffix.casefold() != ".xlsx":
        print("Erro: o arquivo de entrada deve ter extensão .xlsx.", file=sys.stderr)
        return 2
    if not workbook_path.is_file():
        print(f"Erro: arquivo não encontrado: {workbook_path}", file=sys.stderr)
        return 2

    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        sheet = select_sheet(workbook, args.sheet)
        invitations = parse_invitations(sheet)
        sql = build_sql(invitations, workbook_path, sheet.title)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(sql, encoding="utf-8", newline="\n")
    except (OSError, WorkbookValidationError) as error:
        print(f"Erro: {error}", file=sys.stderr)
        return 2

    couples = sum(invitation.invite_type == "couple" for invitation in invitations)
    companions = sum(invitation.max_guests for invitation in invitations)
    print(f"Aba: {sheet.title}")
    print(f"Convites: {len(invitations)} ({couples} de casal)")
    print(f"Acompanhantes: {companions}")
    print(f"SQL: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
